import re
import uuid
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ARTIFACT_DIR = Path(__file__).resolve().parent / "artifacts"
FONT_CANDIDATES = [
    "C:/Windows/Fonts/msyh.ttc",
    "C:/Windows/Fonts/simhei.ttf",
    "C:/Windows/Fonts/simsun.ttc",
]


def normalize_answer_text(text):
    text = re.sub(r"\*\*([^*]+)\*\*", r"\1", text or "")
    text = re.sub(r"\*([^*\n]+)\*", r"\1", text)
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    return "\n".join(line.rstrip() for line in text.split("\n")).strip()


def compact_text(text, limit=120):
    cleaned = re.sub(r"\s+", " ", (text or "").strip())
    if len(cleaned) <= limit:
        return cleaned
    return f"{cleaned[: limit - 1]}..."


def split_label_and_body(text):
    for separator in ("：", ":"):
        if separator in text:
            label, body = text.split(separator, 1)
            return label.strip(), body.strip()
    return text.strip(), ""


def detect_condition(text):
    markers = ("若", "如果", "如", "当", "是否", "确认", "检查", "低于", "大于", "小于", "无法", "仍")
    if any(marker in text for marker in markers):
        return compact_text(text, 80)
    return ""


def detect_devices(text):
    codes = re.findall(r"=\d+-[A-Za-z]\d+", text or "")
    names = []
    for phrase in ("蓄电池", "HMI", "车辆屏", "司控器", "车门", "断路器", "旁路", "受电弓", "高速断路器"):
        if phrase in text:
            names.append(phrase)
    items = codes + [item for item in names if item not in codes]
    return "、".join(dict.fromkeys(items))


def extract_disposal_plan(question, answer, references=None):
    answer_text = normalize_answer_text(answer)
    references = references or []
    lines = [line.strip() for line in answer_text.split("\n") if line.strip()]

    steps = []
    outcomes = []
    safety_notes = []
    devices = []
    current_section = ""

    for line in lines:
        section_label = line.rstrip("：:")
        if re.search(r"(处理结果|安全|注意事项|安全提醒|设备位置|关键注意事项)", section_label):
            current_section = section_label
            continue

        bullet = re.match(r"^[\-•]\s*(.+)$", line)
        numbered = re.match(r"^(\d+)[.、]\s*(.+)$", line)
        chinese_step = re.match(r"^(第[一二三四五六七八九十]+步)[：:]\s*(.+)$", line)

        if "安全" in current_section or "注意" in current_section:
            safety_notes.append(bullet.group(1).strip() if bullet else line)
            continue
        if "处理结果" in current_section:
            outcomes.append(bullet.group(1).strip() if bullet else line)
            continue
        if "设备位置" in current_section or "电气柜" in line or "控制柜" in line or "位于" in line:
            devices.append(line)

        if numbered or chinese_step:
            raw = numbered.group(2) if numbered else chinese_step.group(2)
            phase, body = split_label_and_body(raw)
            full_text = body or phase
            steps.append(
                {
                    "phase": compact_text(phase, 30),
                    "action": compact_text(full_text, 160),
                    "condition": detect_condition(full_text),
                    "device": detect_devices(full_text),
                    "note": "",
                }
            )

    if not steps:
        sentence_candidates = re.split(r"[。；;]\s*", answer_text)
        for sentence in sentence_candidates:
            sentence = sentence.strip()
            if len(sentence) >= 8:
                steps.append(
                    {
                        "phase": f"处置步骤{len(steps) + 1}",
                        "action": compact_text(sentence, 160),
                        "condition": detect_condition(sentence),
                        "device": detect_devices(sentence),
                        "note": "",
                    }
                )
            if len(steps) >= 5:
                break

    if not safety_notes:
        for line in lines:
            if any(marker in line for marker in ("动车前", "必须", "请务必", "注意")):
                safety_notes.append(line)
    if not safety_notes:
        safety_notes.append("动车前确认车门关好、制动缓解，并按运营规程执行。")

    device_text = "\n".join(devices + [step.get("device", "") for step in steps])
    device_codes = re.findall(r"=\d+-[A-Za-z]\d+", device_text)
    device_rows = []
    for code in dict.fromkeys(device_codes):
        location = ""
        for line in devices + lines:
            if code in line and ("位于" in line or "电气柜" in line or "控制柜" in line):
                location = compact_text(line, 120)
                break
        device_rows.append({"code": code, "location": location or "见处置步骤说明"})

    plan = {
        "scenario": compact_text(question, 100),
        "summary": compact_text(lines[0] if lines else answer_text, 180),
        "steps": steps,
        "outcomes": outcomes,
        "safety_notes": safety_notes,
        "devices": device_rows,
        "references": references,
        "generated_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
    }
    return plan


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin_name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin_name}"))
        if node is None:
            node = OxmlElement(f"w:{margin_name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_widths(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row in table.rows:
        for index, width in enumerate(widths):
            if index >= len(row.cells):
                continue
            row.cells[index].width = Inches(width)
            set_cell_margins(row.cells[index])
            row.cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def style_table_header(row):
    for cell in row.cells:
        set_cell_shading(cell, "E8EEF5")
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.bold = True
                run.font.color.rgb = RGBColor(23, 32, 51)


def add_table(document, headers, rows, widths):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    header_cells = table.rows[0].cells
    for index, header in enumerate(headers):
        header_cells[index].text = header
    style_table_header(table.rows[0])

    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = str(value or "")
    set_table_widths(table, widths)
    return table


def setup_document_styles(document):
    section = document.sections[0]
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.85)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

    normal = document.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.space_after = Pt(6)

    for style_name, size, color in (
        ("Title", 22, "0B2545"),
        ("Heading 1", 15, "1F4D78"),
        ("Heading 2", 12.5, "1F4D78"),
    ):
        style = document.styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True


def add_bullet(document, text):
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.add_run(text)
    return paragraph


def generate_docx(plan, output_path, execution_records=None):
    document = Document()
    setup_document_styles(document)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("故障处置建议")

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run(f"轨交车辆故障处置助手 Agent 自动生成 | {plan['generated_at']}")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(83, 98, 118)

    document.add_heading("一、故障问题", level=1)
    document.add_paragraph(plan["scenario"])

    document.add_heading("二、处置建议摘要", level=1)
    document.add_paragraph(plan["summary"])

    document.add_heading("三、处置步骤表", level=1)
    step_rows = []
    for index, step in enumerate(plan["steps"], 1):
        step_rows.append(
            [
                index,
                step.get("phase", ""),
                step.get("action", ""),
                step.get("condition", ""),
                step.get("device", ""),
                step.get("note", ""),
            ]
        )
    add_table(
        document,
        ["序号", "处置阶段", "操作内容", "判断/条件", "涉及设备", "注意事项"],
        step_rows,
        [0.45, 1.0, 2.2, 1.35, 1.0, 1.05],
    )

    if plan["outcomes"]:
        document.add_heading("四、处理结果", level=1)
        for item in plan["outcomes"]:
            add_bullet(document, item)

    document.add_heading("五、设备与开关", level=1)
    device_rows = [[item["code"], item["location"]] for item in plan["devices"]]
    if not device_rows:
        device_rows = [["无明确设备编号", "请按处置步骤和现场规程确认"]]
    add_table(document, ["设备/开关编号", "位置或说明"], device_rows, [1.5, 5.1])

    document.add_heading("六、安全提醒", level=1)
    for item in plan["safety_notes"]:
        add_bullet(document, item)

    document.add_heading("七、Agent执行记录", level=1)
    record_rows = []
    for record in execution_records or []:
        record_rows.append([record.get("stage", ""), record.get("title", ""), record.get("detail", "")])
    if not record_rows:
        record_rows = [["任务规划", "生成故障处置建议", "根据回答提取步骤、设备和安全提醒并生成Word文档"]]
    add_table(document, ["阶段", "动作", "记录"], record_rows, [1.2, 1.8, 3.6])

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    return output_path


def validate_docx(output_path, plan):
    checks = []
    if output_path.exists() and output_path.stat().st_size > 10_000:
        checks.append({"name": "文件生成", "ok": True, "detail": f"文件大小 {output_path.stat().st_size} 字节"})
    else:
        checks.append({"name": "文件生成", "ok": False, "detail": "文件不存在或体积异常"})

    try:
        document = Document(output_path)
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        table_count = len(document.tables)
        checks.append({"name": "可打开性", "ok": True, "detail": "python-docx 可正常打开文档"})
        checks.append({"name": "标题检查", "ok": "故障处置建议" in text, "detail": "检查文档标题"})
        checks.append({"name": "表格检查", "ok": table_count >= 2, "detail": f"检测到 {table_count} 个表格"})
        step_table_rows = len(document.tables[0].rows) - 1 if document.tables else 0
        checks.append(
            {
                "name": "步骤行数",
                "ok": step_table_rows >= max(1, len(plan.get("steps", []))),
                "detail": f"步骤表 {step_table_rows} 行，规划步骤 {len(plan.get('steps', []))} 行",
            }
        )
        checks.append({"name": "安全提醒", "ok": "安全提醒" in text, "detail": "检查安全提醒章节"})
    except Exception as exc:
        checks.append({"name": "可打开性", "ok": False, "detail": str(exc)})

    return {"ok": all(item["ok"] for item in checks), "checks": checks}


def improve_plan_if_needed(plan, validation):
    improvements = []
    if not plan.get("safety_notes"):
        plan["safety_notes"] = ["动车前确认车门关好、制动缓解，并按运营规程执行。"]
        improvements.append("补充默认安全提醒")
    if not plan.get("steps"):
        plan["steps"] = [
            {
                "phase": "人工复核",
                "action": "当前回答未拆解出明确步骤，请结合现场规程人工复核。",
                "condition": "",
                "device": "",
                "note": "",
            }
        ]
        improvements.append("补充人工复核步骤")
    if validation and not validation.get("ok"):
        improvements.append("根据校验结果重新写入完整章节和表格")
    return improvements


def load_font(size, bold=False):
    candidates = ["C:/Windows/Fonts/msyhbd.ttc"] if bold else []
    candidates.extend(FONT_CANDIDATES)
    for candidate in candidates:
        if Path(candidate).exists():
            return ImageFont.truetype(candidate, size)
    return ImageFont.load_default()


def text_size(draw, text, font):
    box = draw.textbbox((0, 0), text, font=font)
    return box[2] - box[0], box[3] - box[1]


def wrap_text(draw, text, font, max_width):
    text = re.sub(r"\s+", " ", text or "").strip()
    if not text:
        return [""]
    lines = []
    current = ""
    for char in text:
        candidate = current + char
        if text_size(draw, candidate, font)[0] <= max_width:
            current = candidate
        else:
            if current:
                lines.append(current)
            current = char
    if current:
        lines.append(current)
    return lines[:5]


def build_flowchart_plan(question, answer, references=None):
    plan = extract_disposal_plan(question, answer, references)
    title = infer_flowchart_title(question, answer)
    nodes = [{"kind": "start", "text": f"发现故障：{infer_fault_name(question, answer)}"}]
    branches = 0

    for step in plan["steps"]:
        action_text = step.get("action") or step.get("phase") or "执行处置步骤"
        nodes.append({"kind": "process", "text": action_text})
        condition = step.get("condition")
        if condition and any(marker in condition for marker in ("若", "如果", "是否", "低于", "大于", "读数", "无法", "仍")):
            nodes.append({"kind": "decision", "text": condition})
            branches += 1

    for outcome in plan.get("outcomes", [])[:2]:
        nodes.append({"kind": "process", "text": outcome})

    if plan.get("safety_notes"):
        safety_note = re.sub(r"^安全提醒[：:]\s*", "", plan["safety_notes"][0])
        nodes.append({"kind": "note", "text": f"安全提醒：{safety_note}"})

    nodes.append({"kind": "end", "text": "完成处置，按调度要求继续运营或退出服务"})

    return {
        "title": title,
        "scenario": plan["scenario"],
        "nodes": nodes,
        "branch_count": branches,
        "safety_count": len(plan.get("safety_notes", [])),
        "generated_at": plan["generated_at"],
    }


def infer_fault_name(question, answer):
    source = f"{question}\n{answer}"
    fault_patterns = [
        ("车门红点故障", ("车门", "红点")),
        ("车门紧急解锁触发紧急制动", ("紧急解锁",)),
        ("紧急制动无法缓解", ("紧急制动", "无法缓解")),
        ("蓄电池无法激活", ("蓄电池", "无法激活")),
        ("高速断路器无法闭合", ("高速断路器", "无法闭合")),
        ("辅助逆变器故障", ("辅助逆变器",)),
        ("牵引系统图标红点故障", ("牵引", "红点")),
    ]
    for name, keywords in fault_patterns:
        if all(keyword in source for keyword in keywords):
            return name

    cleaned = re.sub(r"[？?。；;，,]", " ", question or "").strip()
    cleaned = re.sub(r"\s+", " ", cleaned)
    return compact_text(cleaned, 28) or "车辆故障"


def infer_flowchart_title(question, answer):
    return f"{infer_fault_name(question, answer)}处置流程"


def draw_centered_text(draw, box, lines, font, fill):
    x1, y1, x2, y2 = box
    line_gap = 8
    heights = [text_size(draw, line, font)[1] for line in lines]
    total_height = sum(heights) + line_gap * (len(lines) - 1)
    y = y1 + ((y2 - y1) - total_height) / 2
    for line, height in zip(lines, heights):
        width = text_size(draw, line, font)[0]
        draw.text((x1 + ((x2 - x1) - width) / 2, y), line, font=font, fill=fill)
        y += height + line_gap


def draw_arrow(draw, x, y1, y2, fill="#78909C"):
    draw.line((x, y1, x, y2 - 12), fill=fill, width=4)
    draw.polygon([(x, y2), (x - 10, y2 - 14), (x + 10, y2 - 14)], fill=fill)


def draw_node(draw, node, x, y, width, height, fonts):
    kind = node["kind"]
    palette = {
        "start": ("#E7F3FF", "#2E74B5", "#0B2545"),
        "process": ("#FFFFFF", "#9CB2C7", "#172033"),
        "decision": ("#FFF7E0", "#C8922B", "#172033"),
        "note": ("#FFF4D6", "#D7A53A", "#5C4400"),
        "end": ("#EAF7F1", "#2F8F83", "#0F5F56"),
    }
    fill, outline, text_fill = palette.get(kind, palette["process"])
    box = (x, y, x + width, y + height)
    font = fonts["body"]
    lines = wrap_text(draw, node["text"], font, width - 44)

    if kind == "decision":
        cx = x + width / 2
        cy = y + height / 2
        points = [(cx, y), (x + width, cy), (cx, y + height), (x, cy)]
        draw.polygon(points, fill=fill, outline=outline)
        draw.line(points + [points[0]], fill=outline, width=3)
        draw_centered_text(draw, (x + 28, y + 16, x + width - 28, y + height - 16), lines, font, text_fill)
    else:
        radius = 24 if kind in ("start", "end") else 14
        draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=3)
        draw_centered_text(draw, (x + 20, y + 14, x + width - 20, y + height - 14), lines, font, text_fill)


def generate_flowchart_png(flowchart, output_path):
    width = 1080
    node_width = 760
    node_height = 128
    decision_height = 156
    gap = 76
    header_height = 250
    margin_top = header_height
    margin_bottom = 150
    x = (width - node_width) // 2
    y = margin_top
    heights = [decision_height if node["kind"] == "decision" else node_height for node in flowchart["nodes"]]
    height = margin_top + margin_bottom + sum(heights) + gap * (len(heights) - 1)

    image = Image.new("RGB", (width, height), "#F6F8FB")
    draw = ImageDraw.Draw(image)
    fonts = {
        "title": load_font(36, bold=True),
        "meta": load_font(20),
        "body": load_font(26),
    }

    title_lines = wrap_text(draw, flowchart["title"], fonts["title"], width - 160)
    title_y = 42
    for line in title_lines[:2]:
        line_width = text_size(draw, line, fonts["title"])[0]
        draw.text(((width - line_width) / 2, title_y), line, font=fonts["title"], fill="#0B2545")
        title_y += 44

    scenario = f"场景：{flowchart['scenario']}"
    scenario_lines = wrap_text(draw, scenario, fonts["meta"], width - 180)
    scenario_y = title_y + 10
    for line in scenario_lines[:2]:
        line_width = text_size(draw, line, fonts["meta"])[0]
        draw.text(((width - line_width) / 2, scenario_y), line, font=fonts["meta"], fill="#536276")
        scenario_y += 30

    draw.line((150, header_height - 34, width - 150, header_height - 34), fill="#D6DEE8", width=2)

    for index, node in enumerate(flowchart["nodes"]):
        current_height = heights[index]
        draw_node(draw, node, x, y, node_width, current_height, fonts)
        next_y = y + current_height + gap
        if index < len(flowchart["nodes"]) - 1:
            label = ""
            if node["kind"] == "decision":
                label = "继续处置"
            arrow_start = y + current_height + 10
            arrow_end = next_y - 18
            draw_arrow(draw, width // 2, arrow_start, arrow_end)
            if label:
                draw.rounded_rectangle(
                    (width // 2 + 18, arrow_start + 18, width // 2 + 142, arrow_start + 54),
                    radius=16,
                    fill="#FFFFFF",
                    outline="#D6DEE8",
                    width=2,
                )
                draw.text((width // 2 + 34, arrow_start + 22), label, font=fonts["meta"], fill="#536276")
        y = next_y

    meta = f"轨交车辆故障处置助手 Agent 自动生成 | {flowchart['generated_at']}"
    meta_width = text_size(draw, meta, fonts["meta"])[0]
    draw.text(((width - meta_width) / 2, height - 62), meta, font=fonts["meta"], fill="#8A98AA")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    image.save(output_path, "PNG")
    return output_path


def validate_flowchart_png(output_path, flowchart):
    checks = []
    checks.append(
        {
            "name": "文件生成",
            "ok": output_path.exists() and output_path.stat().st_size > 20_000,
            "detail": f"文件大小 {output_path.stat().st_size if output_path.exists() else 0} 字节",
        }
    )
    try:
        with Image.open(output_path) as image:
            width, height = image.size
            checks.append({"name": "可打开性", "ok": image.format == "PNG", "detail": f"图片格式 {image.format}"})
            checks.append({"name": "画布尺寸", "ok": width >= 900 and height >= 600, "detail": f"尺寸 {width} x {height}"})
            checks.append(
                {
                    "name": "节点数量",
                    "ok": len(flowchart.get("nodes", [])) >= 3,
                    "detail": f"流程节点 {len(flowchart.get('nodes', []))} 个",
                }
            )
            checks.append({"name": "竖版布局", "ok": height > width, "detail": "图片高度大于宽度，适合手机竖向查看"})
    except Exception as exc:
        checks.append({"name": "可打开性", "ok": False, "detail": str(exc)})
    return {"ok": all(item["ok"] for item in checks), "checks": checks}


def improve_flowchart_if_needed(flowchart, validation):
    improvements = []
    if not flowchart.get("nodes"):
        flowchart["nodes"] = [
            {"kind": "start", "text": "开始处置"},
            {"kind": "process", "text": "请结合回答内容和现场规程人工复核"},
            {"kind": "end", "text": "完成处置"},
        ]
        improvements.append("补充基础流程节点")
    if flowchart["nodes"][-1]["kind"] != "end":
        flowchart["nodes"].append({"kind": "end", "text": "完成处置"})
        improvements.append("补充结束节点")
    if not any(node["kind"] == "note" for node in flowchart["nodes"]):
        flowchart["nodes"].insert(-1, {"kind": "note", "text": "安全提醒：动车前确认车门关好、制动缓解。"})
        improvements.append("补充安全提醒节点")
    if validation and not validation.get("ok"):
        improvements.append("根据校验结果重新渲染竖版PNG")
    return improvements


def build_word_artifact(question, answer, references=None):
    ARTIFACT_DIR.mkdir(parents=True, exist_ok=True)
    artifact_id = uuid.uuid4().hex[:10]
    filename = f"故障处置建议-{artifact_id}.docx"
    output_path = ARTIFACT_DIR / filename
    records = []

    def record(stage, title, detail):
        item = {"stage": stage, "title": title, "detail": detail}
        records.append(item)
        return item

    yield {"type": "record", **record("任务规划", "读取上下文", "接收用户问题、模型回答和本轮检索切片")}
    plan = extract_disposal_plan(question, answer, references)
    yield {
        "type": "plan",
        "plan": {
            "scenario": plan["scenario"],
            "step_count": len(plan["steps"]),
            "device_count": len(plan["devices"]),
            "safety_count": len(plan["safety_notes"]),
        },
    }
    yield {
        "type": "record",
        **record(
            "任务规划",
            "拆解处置内容",
            f"识别处置步骤 {len(plan['steps'])} 项，设备/开关 {len(plan['devices'])} 项，安全提醒 {len(plan['safety_notes'])} 项",
        ),
    }
    yield {"type": "record", **record("任务执行", "调用工具 extract_disposal_plan", "将自然语言回答转换为结构化处置计划")}
    yield {"type": "record", **record("任务执行", "调用工具 generate_docx", "写入标题、摘要、处置步骤表、设备表、安全提醒和执行记录")}
    generate_docx(plan, output_path, records)

    yield {"type": "record", **record("自动测试", "调用工具 validate_docx", "检查文件存在、可打开性、标题、表格和步骤行数")}
    validation = validate_docx(output_path, plan)
    for check in validation["checks"]:
        yield {"type": "check", "ok": check["ok"], "title": check["name"], "detail": check["detail"]}

    improvements = improve_plan_if_needed(plan, validation)
    if improvements:
        for improvement in improvements:
            yield {"type": "record", **record("自动改进", "修正文档内容", improvement)}
        generate_docx(plan, output_path, records)
        validation = validate_docx(output_path, plan)
        yield {
            "type": "record",
            **record("自动测试", "复测文档", "重新检查改进后的 Word 文档结构"),
        }
    else:
        yield {"type": "record", **record("自动改进", "无需改进", "自动测试全部通过，保留当前文档")}

    yield {
        "type": "artifact",
        "filename": filename,
        "download_url": f"/api/files/{filename}",
        "validation_ok": validation["ok"],
    }


def build_flowchart_artifact(question, answer, references=None):
    ARTIFACT_DIR.mkdir(parents=True, exist_ok=True)
    artifact_id = uuid.uuid4().hex[:10]
    filename = f"故障处置流程图-{artifact_id}.png"
    output_path = ARTIFACT_DIR / filename
    records = []

    def record(stage, title, detail):
        item = {"stage": stage, "title": title, "detail": detail}
        records.append(item)
        return item

    yield {"type": "record", **record("任务规划", "读取上下文", "接收用户问题、模型回答和本轮检索切片")}
    flowchart = build_flowchart_plan(question, answer, references)
    process_count = sum(1 for node in flowchart["nodes"] if node["kind"] == "process")
    decision_count = sum(1 for node in flowchart["nodes"] if node["kind"] == "decision")
    yield {
        "type": "plan",
        "plan": {
            "scenario": flowchart["scenario"],
            "node_count": len(flowchart["nodes"]),
            "step_count": process_count,
            "branch_count": decision_count,
            "safety_count": flowchart["safety_count"],
        },
    }
    yield {
        "type": "record",
        **record(
            "任务规划",
            "规划竖版流程图",
            f"识别流程节点 {len(flowchart['nodes'])} 个，操作节点 {process_count} 个，判断节点 {decision_count} 个",
        ),
    }
    yield {"type": "record", **record("任务执行", "调用工具 extract_flowchart_plan", "将自然语言回答转换为流程图结构")}
    yield {"type": "record", **record("任务执行", "调用工具 generate_flowchart_png", "绘制标题、节点、箭头、判断提示和安全提醒")}
    generate_flowchart_png(flowchart, output_path)

    yield {"type": "record", **record("自动测试", "调用工具 validate_flowchart_png", "检查PNG文件、可打开性、画布尺寸、节点数量和竖版布局")}
    validation = validate_flowchart_png(output_path, flowchart)
    for check in validation["checks"]:
        yield {"type": "check", "ok": check["ok"], "title": check["name"], "detail": check["detail"]}

    improvements = improve_flowchart_if_needed(flowchart, validation)
    if improvements:
        for improvement in improvements:
            yield {"type": "record", **record("自动改进", "修正流程图结构", improvement)}
        generate_flowchart_png(flowchart, output_path)
        validation = validate_flowchart_png(output_path, flowchart)
        yield {"type": "record", **record("自动测试", "复测流程图", "重新检查改进后的 PNG 流程图")}
    else:
        yield {"type": "record", **record("自动改进", "无需改进", "自动测试全部通过，保留当前流程图")}

    yield {
        "type": "artifact",
        "filename": filename,
        "download_url": f"/api/files/{filename}",
        "preview_url": f"/api/files/{filename}?preview=1",
        "validation_ok": validation["ok"],
    }
